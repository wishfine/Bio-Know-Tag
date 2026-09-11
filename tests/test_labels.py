from pathlib import Path

import pytest
from openpyxl import Workbook

from bio_know_tag.labels import load_label_workbook, write_label_outputs


HEADERS = (
    "全路径知识点名称",
    "末级知识点编号",
    "知识点名称",
    "知识点类型",
    "定义 / 核心内容",
    "核心概念 / 关键过程",
    "常见考查方式",
    "易混淆区分",
)


def make_workbook(
    tmp_path: Path,
    *,
    label_id: str = "001",
    label_name: str = "观察实验（旧）\u00a0",
    names: list[str] | None = None,
) -> Path:
    path = tmp_path / "labels.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "知识点及试题量详情"
    sheet.append(HEADERS)
    for index, name in enumerate(names or [label_name], start=1):
        sheet.append(
            (
                f"知识点 -> 模块 -> {name}",
                label_id if index == 1 else f"{index:03d}",
                name,
                "知识",
                " 定义\n内容 ",
                "核心\u3000过程",
                "选择题",
                "相近内容",
            )
        )
    workbook.save(path)
    return path


def test_load_label_workbook_normalizes_whitespace_and_preserves_id(tmp_path: Path):
    workbook = make_workbook(tmp_path)

    labels = load_label_workbook(workbook)

    assert labels == [
        {
            "label_path": "知识点 -> 模块 -> 观察实验（旧）",
            "label_id": "001",
            "label_name": "观察实验（旧）",
            "label_type": "知识",
            "definition": "定义 内容",
            "core_concepts": "核心 过程",
            "common_assessments": "选择题",
            "distinctions": "相近内容",
        }
    ]


def test_load_label_workbook_rejects_duplicate_names(tmp_path: Path):
    workbook = make_workbook(tmp_path, names=["体液免疫", "体液免疫"])

    with pytest.raises(ValueError, match="duplicate label_name: 体液免疫"):
        load_label_workbook(workbook)


def test_load_label_workbook_rejects_missing_required_value(tmp_path: Path):
    workbook = make_workbook(tmp_path)
    from openpyxl import load_workbook

    editable = load_workbook(workbook)
    editable.active["E2"] = ""
    editable.save(workbook)

    with pytest.raises(ValueError, match="row 2.*definition"):
        load_label_workbook(workbook)


def test_write_label_outputs_writes_jsonl_and_quality_report(tmp_path: Path):
    records = load_label_workbook(make_workbook(tmp_path, label_name="光合作用综合"))
    output_path = tmp_path / "labels.jsonl"
    report_path = tmp_path / "labels.report.json"

    report = write_label_outputs(records, output_path, report_path)

    assert output_path.read_text(encoding="utf-8").count("\n") == 1
    assert report == {
        "input": 1,
        "processed": 1,
        "error": 0,
        "blank_fields": 0,
        "duplicate_ids": 0,
        "duplicate_names": 0,
        "category_counts": {"其他": 0, "实验": 0, "应用": 0, "综合": 1},
    }
    assert '"processed": 1' in report_path.read_text(encoding="utf-8")

