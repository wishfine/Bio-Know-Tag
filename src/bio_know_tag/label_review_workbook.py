"""Export the per-label review ledger to a filterable Excel workbook."""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


SUMMARY_HEADERS = [
    "序号",
    "label_id",
    "Label名称",
    "Label路径",
    "原释义",
    "DS释义",
    "DS Judge结果",
    "GPT Judge结果（二次复核）",
    "处理策略",
    "人工跟进",
    "风险/备注",
]

DETAIL_HEADERS = [
    "序号",
    "label_id",
    "Label名称",
    "Label路径",
    "Label类型",
    "原定义",
    "原核心概念",
    "原常见考查",
    "原易混淆区分",
    "DS核心含义",
    "DS包含内容",
    "DS排除内容",
    "DS Judge类别",
    "DS Judge分数",
    "DS Judge名称充分性",
    "DS Judge结论",
    "DS Judge审核理由",
    "DS Judge边界差异",
    "DS Judge遗漏",
    "DS Judge扩展",
    "GPT Judge状态",
    "GPT Judge置信度",
    "GPT Judge此前策略",
    "GPT Judge最终策略",
    "GPT Judge是否调整",
    "GPT Judge复核结论",
    "GPT Judge人工跟进",
    "GPT Judge核对证据",
    "业务边界覆盖",
    "最终处理策略",
    "策略说明",
    "自动化方式",
    "Prompt字段",
    "需要老师释义",
    "策略置信度",
    "已知风险",
    "问题分类",
    "taxonomy风险/问题",
    "参考策略代码",
    "参考策略提示",
]

FOLLOWUP_HEADERS = [
    "序号",
    "label_id",
    "Label名称",
    "Label路径",
    "问题分类",
    "DS Judge类别",
    "DS Judge分数",
    "GPT Judge状态",
    "最终处理策略",
    "人工跟进原因",
]

TEACHER_REVIEW_HEADERS = [
    "序号",
    "label_id",
    "Label名称",
    "Label路径",
    "原释义（老师）",
    "DS释义",
    "DS Judge结果",
    "请老师重点确认",
    "老师复核结论",
    "老师修改建议",
]

TEACHER_CONFIRMATIONS = {
    "酶的特性": "是否将“作用条件较温和”纳入本Label；温度、pH影响是否归入本Label；“多样性”是否继续作为独立特性。",
    "基于代谢类型对生物进行分类": "是否严格按碳源×能源划分四类；需氧型、厌氧型、兼性厌氧型是否排除；光能异养型是否属于高中考查范围。",
    "光合作用综合": "是否包含化能合成作用、色素提取与分离实验及生态层面意义；满足什么条件才打“综合”（当前规则为覆盖至少4个子模块）。",
    "其余伴性遗传疾病": "请确认疾病白名单：是否只排除红绿色盲、明确包含血友病；白名单之外的伴性遗传病是否命中。",
    "RNA分子的种类与功能": "是否包含核酶；RNA与DNA的结构比较是否归入本Label；真核mRNA的5'帽和poly(A)尾是否超出范围。",
    "蛋白质病毒的增殖": "Label名称与释义明显不一致。请确认是否改名为“RNA病毒及逆转录病毒的增殖”；若保留现名，请明确“蛋白质病毒”的含义。",
    "遗传与变异综合": "是否限定为变异—遗传病—育种—进化的综合；孟德尔遗传规律和减数分裂是否纳入；满足什么条件才打“综合”。",
    "生态系统的营养结构": "是否包含食物网复杂程度与生态系统稳定性的关系；数量、生物量和能量金字塔是否属于本Label。",
    "蛋白质工程的进程与前景": "是否需要纳入蛋白质工程的基本定义与操作流程，还是严格限定为发展进程和应用前景。",
    "基因工程综合": "是否包含蛋白质工程；安全性与伦理问题是否纳入；满足什么条件才打综合Label而不是单个子Label。",
}


def _plain(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, (list, tuple)):
        return "\n".join(f"{index}. {item}" for index, item in enumerate(value, 1))
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value)


def _section(title: str, value: Any) -> str:
    rendered = _plain(value)
    return f"{title}：{rendered}" if rendered else ""


def _join_sections(*sections: str) -> str:
    return "\n".join(section for section in sections if section)


def _teacher_definition(record: dict[str, Any]) -> str:
    return _join_sections(
        _section("定义", record.get("definition")),
        _section("核心概念", record.get("core_concepts")),
        _section("常见考查", record.get("common_assessments")),
        _section("易混淆区分", record.get("distinctions")),
    )


def _ds_definition(record: dict[str, Any]) -> str:
    stage1 = record.get("stage1") or {}
    return _join_sections(
        _section("核心含义", stage1.get("core_meaning")),
        _section("包含内容", stage1.get("included_content")),
        _section("排除内容", stage1.get("excluded_content")),
    )


def _ds_judge(record: dict[str, Any]) -> str:
    judge = record.get("stage2_judge") or {}
    return _join_sections(
        _section("类别", record.get("stage2_category")),
        _section("分数", judge.get("alignment_score")),
        _section("名称充分性", judge.get("name_sufficiency")),
        _section("结论", judge.get("audit_decision")),
        _section("审核理由", judge.get("audit_reason")),
        _section("边界差异", judge.get("boundary_differences")),
        _section("遗漏", judge.get("omissions")),
        _section("扩展", judge.get("expansions")),
    )


def _gpt_judge(record: dict[str, Any]) -> str:
    review = record.get("second_review") or {}
    final_strategy = record.get("final_strategy") or {}
    return _join_sections(
        _section("状态", review.get("status")),
        _section("置信度", review.get("confidence")),
        _section("此前策略", review.get("previous_mode")),
        _section("最终策略", review.get("final_mode") or final_strategy.get("mode")),
        _section("是否调整", review.get("adjusted_from_previous")),
        _section("复核结论", review.get("rationale")),
        _section("人工跟进", review.get("manual_followup_required")),
    )


def _strategy(record: dict[str, Any]) -> str:
    final_strategy = record.get("final_strategy") or {}
    previous_strategy = record.get("strategy") or {}
    review = record.get("second_review") or {}
    return _join_sections(
        _section("模式", final_strategy.get("mode")),
        _section("说明", final_strategy.get("reason") or previous_strategy.get("reason") or review.get("rationale")),
        _section("自动化", final_strategy.get("automation") or previous_strategy.get("automation")),
        _section("Prompt字段", final_strategy.get("prompt_fields") or previous_strategy.get("prompt_fields")),
        _section("业务边界", final_strategy.get("operational_boundary") or review.get("operational_boundary")),
        _section(
            "需要老师释义",
            final_strategy.get("teacher_definition_required")
            if "teacher_definition_required" in final_strategy
            else previous_strategy.get("teacher_definition_required"),
        ),
        _section("置信度", final_strategy.get("confidence") or previous_strategy.get("confidence") or review.get("confidence")),
    )


def _issue_type(record: dict[str, Any]) -> str:
    review = record.get("second_review") or {}
    issue = record.get("taxonomy_issue") or {}
    name = str(record.get("label_name", ""))
    if review.get("final_mode") == "taxonomy_hold":
        if name == "蛋白质病毒的增殖":
            return "Label名称与原释义冲突"
        return "taxonomy冲突"
    if review.get("operational_boundary"):
        return "业务侧重点边界覆盖"
    if issue.get("风险级别") in {"P1", "P2"}:
        return "Label范围重叠或边界开放"
    judge = record.get("stage2_judge") or {}
    if judge.get("alignment_score", 5) <= 3 or judge.get("audit_decision") != "两者基本等价":
        return "DS与老师原释义存在明显差异"
    if review.get("status") == "adjusted":
        return "二次复核发现关键措辞风险"
    return "—"


def _detail_row(index: int, record: dict[str, Any]) -> dict[str, Any]:
    stage1 = record.get("stage1") or {}
    judge = record.get("stage2_judge") or {}
    review = record.get("second_review") or {}
    final_strategy = record.get("final_strategy") or {}
    previous_strategy = record.get("strategy") or {}
    reference = record.get("reference_strategy") or {}
    issue = record.get("taxonomy_issue") or {}
    return {
        "序号": index,
        "label_id": record.get("label_id", ""),
        "Label名称": record.get("label_name", ""),
        "Label路径": record.get("label_path", ""),
        "Label类型": record.get("label_type", ""),
        "原定义": record.get("definition", ""),
        "原核心概念": record.get("core_concepts", ""),
        "原常见考查": record.get("common_assessments", ""),
        "原易混淆区分": record.get("distinctions", ""),
        "DS核心含义": stage1.get("core_meaning", ""),
        "DS包含内容": stage1.get("included_content", []),
        "DS排除内容": stage1.get("excluded_content", []),
        "DS Judge类别": record.get("stage2_category", ""),
        "DS Judge分数": judge.get("alignment_score", ""),
        "DS Judge名称充分性": judge.get("name_sufficiency", ""),
        "DS Judge结论": judge.get("audit_decision", ""),
        "DS Judge审核理由": judge.get("audit_reason", ""),
        "DS Judge边界差异": judge.get("boundary_differences", []),
        "DS Judge遗漏": judge.get("omissions", []),
        "DS Judge扩展": judge.get("expansions", []),
        "GPT Judge状态": review.get("status", ""),
        "GPT Judge置信度": review.get("confidence", ""),
        "GPT Judge此前策略": review.get("previous_mode", ""),
        "GPT Judge最终策略": review.get("final_mode", ""),
        "GPT Judge是否调整": review.get("adjusted_from_previous", ""),
        "GPT Judge复核结论": review.get("rationale", ""),
        "GPT Judge人工跟进": review.get("manual_followup_required", ""),
        "GPT Judge核对证据": review.get("evidence_checked", []),
        "业务边界覆盖": review.get("operational_boundary", ""),
        "最终处理策略": final_strategy.get("mode", ""),
        "策略说明": final_strategy.get("reason", "") or previous_strategy.get("reason", "") or review.get("rationale", ""),
        "自动化方式": final_strategy.get("automation", "") or previous_strategy.get("automation", ""),
        "Prompt字段": final_strategy.get("prompt_fields", []) or previous_strategy.get("prompt_fields", []),
        "需要老师释义": final_strategy.get("teacher_definition_required", "")
        if "teacher_definition_required" in final_strategy
        else previous_strategy.get("teacher_definition_required", ""),
        "策略置信度": final_strategy.get("confidence", "") or previous_strategy.get("confidence", "") or review.get("confidence", ""),
        "已知风险": final_strategy.get("known_risk_level", "") or review.get("known_risk_level", "") or previous_strategy.get("known_risk_level", ""),
        "问题分类": _issue_type(record),
        "taxonomy风险/问题": _join_sections(
            _section("风险级别", issue.get("风险级别")),
            _section("问题", issue.get("问题")),
            _section("建议", issue.get("建议")),
        ),
        "参考策略代码": reference.get("关键词策略代码", ""),
        "参考策略提示": reference.get("Prompt是否需要释义", ""),
    }


def build_workbook_rows(records: Iterable[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    source_records = list(records)
    detail = [_detail_row(index, record) for index, record in enumerate(source_records, 1)]
    summary = [
        {
            "序号": row["序号"],
            "label_id": row["label_id"],
            "Label名称": row["Label名称"],
            "Label路径": row["Label路径"],
            "原释义": _join_sections(
                _section("定义", row["原定义"]),
                _section("核心概念", row["原核心概念"]),
                _section("常见考查", row["原常见考查"]),
                _section("易混淆区分", row["原易混淆区分"]),
            ),
            "DS释义": _join_sections(
                _section("核心含义", row["DS核心含义"]),
                _section("包含内容", row["DS包含内容"]),
                _section("排除内容", row["DS排除内容"]),
            ),
            "DS Judge结果": _join_sections(
                _section("类别", row["DS Judge类别"]),
                _section("分数", row["DS Judge分数"]),
                _section("名称充分性", row["DS Judge名称充分性"]),
                _section("结论", row["DS Judge结论"]),
                _section("理由", row["DS Judge审核理由"]),
                _section("边界差异", row["DS Judge边界差异"]),
                _section("遗漏", row["DS Judge遗漏"]),
                _section("扩展", row["DS Judge扩展"]),
            ),
            "GPT Judge结果（二次复核）": _join_sections(
                _section("状态", row["GPT Judge状态"]),
                _section("置信度", row["GPT Judge置信度"]),
                _section("此前策略", row["GPT Judge此前策略"]),
                _section("最终策略", row["GPT Judge最终策略"]),
                _section("是否调整", row["GPT Judge是否调整"]),
                _section("复核结论", row["GPT Judge复核结论"]),
                _section("人工跟进", row["GPT Judge人工跟进"]),
            ),
            "处理策略": _strategy({"final_strategy": {"mode": row["最终处理策略"], "automation": row["自动化方式"], "prompt_fields": row["Prompt字段"], "confidence": row["策略置信度"]}, "strategy": {"reason": row["策略说明"]}, "second_review": {"operational_boundary": row["业务边界覆盖"]}, "label_name": row["Label名称"]}),
            "人工跟进": row["GPT Judge人工跟进"],
            "风险/备注": _join_sections(
                _section("问题分类", row["问题分类"]),
                _section("已知风险", row["已知风险"]),
                _section("taxonomy", row["taxonomy风险/问题"]),
            ),
        }
        for row in detail
    ]
    manual = [
        {
            "序号": index,
            "label_id": row["label_id"],
            "Label名称": row["Label名称"],
            "Label路径": row["Label路径"],
            "问题分类": row["问题分类"],
            "DS Judge类别": row["DS Judge类别"],
            "DS Judge分数": row["DS Judge分数"],
            "GPT Judge状态": row["GPT Judge状态"],
            "最终处理策略": row["最终处理策略"],
            "人工跟进原因": row["GPT Judge复核结论"] or row["taxonomy风险/问题"],
        }
        for index, row in enumerate(
            (row for row in detail if row["GPT Judge人工跟进"] is True),
            1,
        )
    ]
    source_by_id = {str(record.get("label_id", "")): record for record in source_records}
    teacher_review = []
    for row in detail:
        record = source_by_id[row["label_id"]]
        judge = record.get("stage2_judge") or {}
        name = row["Label名称"]
        if judge.get("audit_decision") != "两者都有问题" and name != "蛋白质病毒的增殖":
            continue
        teacher_review.append(
            {
                "序号": len(teacher_review) + 1,
                "label_id": row["label_id"],
                "Label名称": name,
                "Label路径": row["Label路径"],
                "原释义（老师）": _teacher_definition(record),
                "DS释义": _ds_definition(record),
                "DS Judge结果": _ds_judge(record),
                "请老师重点确认": TEACHER_CONFIRMATIONS[name],
                "老师复核结论": "",
                "老师修改建议": "",
            }
        )
    return {"summary": summary, "detail": detail, "manual": manual, "teacher_review": teacher_review}


def _as_cell(value: Any) -> Any:
    return _plain(value) if isinstance(value, (list, tuple, dict)) else value


def _header_column(headers: list[str], header: str) -> str:
    index = headers.index(header) + 1
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _write_table_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[dict[str, Any]], table_name: str, widths: dict[str, int]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.append(headers)
    for row in rows:
        sheet.append([_as_cell(row.get(header, "")) for header in headers])
    for cell in sheet[1]:
        value = str(cell.value or "")
        if value.startswith("原"):
            fill = "548235"
        elif value.startswith("DS"):
            fill = "1F4E78"
        elif value.startswith("GPT"):
            fill = "7030A0"
        elif value in {"处理策略", "策略说明", "自动化方式", "Prompt字段", "最终处理策略"}:
            fill = "C65911"
        elif value in {"问题分类", "风险/备注", "人工跟进", "人工跟进原因"}:
            fill = "BF9000"
        else:
            fill = "44546A"
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    border = Border(bottom=Side(style="thin", color="D9E1F2"))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
        max_lines = max((str(cell.value or "").count("\n") + 1 for cell in row), default=1)
        sheet.row_dimensions[row[0].row].height = min(180, max(24, 15 * max_lines))
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = widths.get(header, 36)
    sheet.auto_filter.ref = f"A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
    if rows:
        table = Table(displayName=table_name, ref=f"A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    return sheet


def _write_summary_sheet(workbook: Workbook, detail_rows: list[dict[str, Any]], summary_rows: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet("统计汇总")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(["指标", "公式结果", "本地核对值"])
    last = len(summary_rows) + 1
    sheet.append(["Label总数", f"=COUNTA('逐Label总表'!$C$2:$C${last})", len(summary_rows)])

    def add_counts(title: str, header: str, values: list[str], counter: Counter[str]) -> None:
        sheet.append([title, "", ""])
        column = _header_column(DETAIL_HEADERS, header)
        for value in values:
            sheet.append([f"{header}={value}", f'=COUNTIF(\'逐Label详情\'!${column}$2:${column}${last},"{value}")', counter.get(value, 0)])

    add_counts("DS Judge类别分布", "DS Judge类别", ["L1", "L2", "L3"], Counter(str(r["DS Judge类别"]) for r in detail_rows))
    add_counts("DS Judge分数分布", "DS Judge分数", ["1", "2", "3", "4", "5"], Counter(str(r["DS Judge分数"]) for r in detail_rows))
    add_counts("GPT Judge状态分布", "GPT Judge状态", ["confirmed", "confirmed_with_caution", "adjusted", "routed_separately", "taxonomy_hold"], Counter(str(r["GPT Judge状态"]) for r in detail_rows))
    add_counts("最终处理策略分布", "最终处理策略", ["name_only", "name_plus_boundary", "compact_definition", "strict_definition", "separate_dimension", "taxonomy_hold"], Counter(str(r["最终处理策略"]) for r in detail_rows))
    add_counts("问题分类分布", "问题分类", ["DS与老师原释义存在明显差异", "Label范围重叠或边界开放", "二次复核发现关键措辞风险", "业务侧重点边界覆盖", "Label名称与原释义冲突", "taxonomy冲突", "—"], Counter(str(r["问题分类"]) for r in detail_rows))
    manual_column = _header_column(DETAIL_HEADERS, "GPT Judge人工跟进")
    adjusted_column = _header_column(DETAIL_HEADERS, "GPT Judge是否调整")
    sheet.append(["需要人工跟进的Label", f"=COUNTIF('逐Label详情'!${manual_column}$2:${manual_column}${last},TRUE)", sum(r["GPT Judge人工跟进"] is True for r in detail_rows)])
    sheet.append(["二次复核调整数", f"=COUNTIF('逐Label详情'!${adjusted_column}$2:${adjusted_column}${last},TRUE)", sum(r["GPT Judge是否调整"] is True for r in detail_rows)])
    sheet.append([])
    sheet.append(["说明", "公式结果在 Excel/WPS 打开时自动重算；本地核对值由导出脚本计算。", ""])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="0F6B78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 62
    sheet.column_dimensions["C"].width = 16


def _write_guide_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("字段说明")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(["字段", "含义", "来源/用途"])
    rows = [
        ("原释义", "老师图谱的 definition、core_concepts、common_assessments、distinctions。", "configs/labels.jsonl"),
        ("DS释义", "Stage1 仅根据 Label 名称生成的核心含义、包含内容和排除内容。", "Stage1 evidence"),
        ("DS Judge", "Stage2 对老师原释义和 DS 释义的对齐评分、结论、遗漏、扩张及边界差异。", "Stage2 evidence"),
        ("GPT Judge（二次复核）", "second_review 的策略复核结果；当前是证据驱动的规则复核，不是另一次 DS API 调用。", "configs/label_strategies.review2.jsonl"),
        ("最终处理策略", "最终用于题目标注的路由。", "final_strategy"),
        ("问题分类", "把高风险行按 DS 差异、范围重叠、措辞风险、业务边界覆盖和名称/释义冲突分类。", "导出脚本"),
        ("人工跟进", "只列需要人工确认或治理的 Label，不代表整条数据处理失败。", "second_review.manual_followup_required"),
        ("老师复核10项", "单列9个“两者都有问题”的Label和“蛋白质病毒的增殖”，预留老师结论与修改建议。", "Stage1 + Stage2 evidence"),
        ("数据范围", "本工作簿由 458 条 review2 台账记录导出。", "repo ledger"),
    ]
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="44546A")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 78
    sheet.column_dimensions["C"].width = 38


def write_workbook(records: Iterable[dict[str, Any]], output_path: str | Path) -> dict[str, Any]:
    source_records = list(records)
    rows = build_workbook_rows(source_records)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary_widths = {"序号": 8, "label_id": 23, "Label名称": 26, "Label路径": 48, "原释义": 60, "DS释义": 60, "DS Judge结果": 60, "GPT Judge结果（二次复核）": 60, "处理策略": 52, "人工跟进": 12, "风险/备注": 52}
    detail_widths = {"序号": 8, "label_id": 23, "Label名称": 26, "Label路径": 48, "Label类型": 10, "DS Judge分数": 12, "DS Judge类别": 12, "GPT Judge状态": 18, "GPT Judge置信度": 12, "最终处理策略": 22, "问题分类": 28, "Prompt字段": 30}
    # Reuse the default sheet for the summary so it remains the first tab.
    default = workbook.active
    default.title = "逐Label总表"
    default.sheet_view.showGridLines = False
    default.freeze_panes = "A2"
    default.append(SUMMARY_HEADERS)
    for row in rows["summary"]:
        default.append([_as_cell(row.get(header, "")) for header in SUMMARY_HEADERS])
    # Style and table for the reused sheet.
    _style_table_sheet(default, SUMMARY_HEADERS, rows["summary"], "LabelReviewSummary", summary_widths)
    _write_table_sheet(workbook, "逐Label详情", DETAIL_HEADERS, rows["detail"], "LabelReviewDetail", detail_widths)
    followup_widths = {"序号": 8, "label_id": 23, "Label名称": 30, "Label路径": 48, "问题分类": 28, "DS Judge类别": 12, "DS Judge分数": 10, "GPT Judge状态": 20, "最终处理策略": 22, "人工跟进原因": 72}
    _write_table_sheet(workbook, "人工跟进", FOLLOWUP_HEADERS, rows["manual"], "LabelReviewFollowup", followup_widths)
    teacher_review_widths = {
        "序号": 8,
        "label_id": 23,
        "Label名称": 30,
        "Label路径": 48,
        "原释义（老师）": 72,
        "DS释义": 72,
        "DS Judge结果": 72,
        "请老师重点确认": 58,
        "老师复核结论": 28,
        "老师修改建议": 58,
    }
    teacher_sheet = _write_table_sheet(
        workbook,
        "老师复核10项",
        TEACHER_REVIEW_HEADERS,
        rows["teacher_review"],
        "LabelTeacherReview",
        teacher_review_widths,
    )
    teacher_sheet.freeze_panes = "E2"
    teacher_sheet.sheet_view.zoomScale = 70
    for row_index in range(2, teacher_sheet.max_row + 1):
        teacher_sheet.row_dimensions[row_index].height = 260
        for column_index in (9, 10):
            teacher_sheet.cell(row_index, column_index).fill = PatternFill("solid", fgColor="DDEBF7")
    _write_summary_sheet(workbook, rows["detail"], rows["summary"])
    _write_guide_sheet(workbook)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output)
    return {
        "output": str(output),
        "rows": len(source_records),
        "manual_followup_rows": len(rows["manual"]),
        "teacher_review_rows": len(rows["teacher_review"]),
        "sheets": workbook.sheetnames,
    }


def _style_table_sheet(sheet, headers: list[str], rows: list[dict[str, Any]], table_name: str, widths: dict[str, int]) -> None:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    for cell in sheet[1]:
        value = str(cell.value or "")
        if value.startswith("原"):
            fill = "548235"
        elif value.startswith("DS"):
            fill = "1F4E78"
        elif value.startswith("GPT"):
            fill = "7030A0"
        elif value in {"处理策略", "策略说明", "自动化方式", "Prompt字段", "最终处理策略"}:
            fill = "C65911"
        elif value in {"问题分类", "风险/备注", "人工跟进", "人工跟进原因"}:
            fill = "BF9000"
        else:
            fill = "44546A"
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    border = Border(bottom=Side(style="thin", color="D9E1F2"))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
        max_lines = max((str(cell.value or "").count("\n") + 1 for cell in row), default=1)
        sheet.row_dimensions[row[0].row].height = min(180, max(24, 15 * max_lines))
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = widths.get(header, 36)
    sheet.auto_filter.ref = f"A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
    if rows:
        table = Table(displayName=table_name, ref=f"A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)


def read_jsonl(path: str | Path) -> list[dict[str, Any]]:
    with Path(path).open(encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]
