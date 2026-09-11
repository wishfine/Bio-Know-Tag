"""Load and validate the biology label taxonomy workbook."""

from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


COLUMN_MAP = {
    "全路径知识点名称": "label_path",
    "末级知识点编号": "label_id",
    "知识点名称": "label_name",
    "知识点类型": "label_type",
    "定义 / 核心内容": "definition",
    "核心概念 / 关键过程": "core_concepts",
    "常见考查方式": "common_assessments",
    "易混淆区分": "distinctions",
}
REQUIRED_FIELDS = tuple(COLUMN_MAP.values())


def normalize_cell(value: Any) -> str:
    """Return a stable, single-line string without altering identifier digits."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value).replace("\u3000", " ")).strip()


def _assert_unique(records: list[dict[str, str]], field: str) -> None:
    counts = Counter(record[field] for record in records)
    duplicates = sorted(value for value, count in counts.items() if count > 1)
    if duplicates:
        raise ValueError(f"duplicate {field}: {duplicates[0]}")


def load_label_workbook(path: str | Path) -> list[dict[str, str]]:
    """Load the first worksheet and enforce the observed taxonomy contract."""
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [normalize_cell(value) for value in next(rows)]
    except StopIteration as exc:
        raise ValueError("workbook is empty") from exc

    missing_columns = [column for column in COLUMN_MAP if column not in headers]
    if missing_columns:
        raise ValueError(f"missing columns: {', '.join(missing_columns)}")
    indexes = {target: headers.index(source) for source, target in COLUMN_MAP.items()}

    records: list[dict[str, str]] = []
    for row_number, row in enumerate(rows, start=2):
        if not any(normalize_cell(value) for value in row):
            continue
        record = {
            field: normalize_cell(row[index] if index < len(row) else None)
            for field, index in indexes.items()
        }
        for field in REQUIRED_FIELDS:
            if not record[field]:
                raise ValueError(f"row {row_number} has blank {field}")
        records.append(record)

    _assert_unique(records, "label_id")
    _assert_unique(records, "label_name")
    _assert_unique(records, "label_path")
    return records


def _atomic_write_lines(path: Path, lines: Iterable[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    with temporary.open("w", encoding="utf-8", newline="\n") as handle:
        for line in lines:
            handle.write(line)
            handle.write("\n")
    temporary.replace(path)


def write_label_outputs(
    records: list[dict[str, str]],
    output_path: str | Path,
    report_path: str | Path,
) -> dict[str, Any]:
    """Write normalized taxonomy JSONL and a deterministic quality report."""
    output = Path(output_path)
    report_output = Path(report_path)
    _atomic_write_lines(
        output,
        (json.dumps(record, ensure_ascii=False, sort_keys=True) for record in records),
    )
    report = {
        "input": len(records),
        "processed": len(records),
        "error": 0,
        "blank_fields": sum(
            1 for record in records for field in REQUIRED_FIELDS if not record.get(field)
        ),
        "duplicate_ids": len(records) - len({record["label_id"] for record in records}),
        "duplicate_names": len(records)
        - len({record["label_name"] for record in records}),
        "category_counts": {
            category: sum(category in record["label_name"] for record in records)
            for category in ("其他", "实验", "应用", "综合")
        },
    }
    _atomic_write_lines(
        report_output,
        [json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True)],
    )
    return report

